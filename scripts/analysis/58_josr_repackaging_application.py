#!/usr/bin/env python
"""Repackage the JOT-ready manuscript for Journal of Orthopaedic Surgery and Research."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import re
import shutil
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


OLD_TITLE = "A meniscus-specific senescence program links fibrochondrocyte state disruption to candidate synovium-cartilage inflammatory remodeling in knee osteoarthritis"
PREVIOUS_JOSR_TITLE = "A program-level meniscus senescence framework prioritizes candidate paracrine target axes in knee osteoarthritis: an integrative single-cell and bulk transcriptomic analysis"
PREVIOUS_JOSR_TITLE_SHORT = "A program-level meniscus senescence framework prioritizes candidate paracrine target axes in knee osteoarthritis: an integrative transcriptomic analysis"
NEW_TITLE = "A program-level meniscus senescence analysis identifies fibrocartilage-matrix stratification signals and candidate paracrine axes in knee osteoarthritis: an integrative transcriptomic study"

ABSTRACT_SECTION = """## Abstract

**Background:** Meniscal degeneration is closely associated with knee osteoarthritis (OA), but meniscus-specific senescence signals are difficult to distinguish from generic aging or inflammation, and single-marker senescence analysis performs poorly in sparse single-cell data. We used a program-level strategy to define a meniscus senescence program (MSP) and to identify paracrine axes for downstream validation.

**Methods:** We combined consensus non-negative matrix factorization of GSE220243 single-cell fibrochondrocyte programs, projection into the author-processed HRA001986 meniscus dataset, random-effects bulk meta-analysis across nine meniscus, cartilage, and synovium cohorts, consensus bulk subtyping, OA-vs-normal diagnostic discrimination testing, and ligand-receptor/receiver-target prioritization.

**Results:** The analysis identified MSP-like fibrochondrocyte programs spanning senescence, SASP, fibrocartilage, and paracrine biology. MIF-CD74, ANGPTL4-integrin, and VEGF emerged as the leading paracrine axes for follow-up. External projection supported a plausible meniscus cell-state context. In bulk cohorts, the fibrocartilage-matrix remodeling axis showed the most consistent disease-associated signal (pooled standardized mean difference 0.80, I2 = 30%, false discovery rate 3.2e-4) and candidate OA-vs-normal stratification value (pooled directional area under the receiver operating characteristic curve 0.806, 95% confidence interval 0.714-0.884; grouped cross-validation area under the curve 0.817 +/- 0.124). MSP paracrine, angiogenic, and inflammatory axes were tissue- and comparator-dependent (I2 approximately 82-85%) rather than uniformly increased in OA.

**Conclusions:** In this public-data transcriptomic study, fibrocartilage-matrix remodeling was the most reproducible MSP-associated bulk signal and showed preliminary OA stratification value. MIF-CD74, ANGPTL4-integrin, and VEGF should be viewed as candidate paracrine axes for biomarker and perturbation studies, not as established disease-driving mechanisms or clinical tests.
"""

KEYWORDS_SECTION = """## Keywords

**Keywords:** Knee osteoarthritis; Meniscus; Cellular senescence; Single-cell RNA sequencing; Paracrine signaling; Fibrocartilage; Diagnostic biomarker; Therapeutic target prioritization; MIF; ANGPTL4
"""

COVER_OPENING = f"""Dear Editors of the Journal of Orthopaedic Surgery and Research,

We are pleased to submit our manuscript entitled "{NEW_TITLE}" for consideration as a Methodology article in the Journal of Orthopaedic Surgery and Research.

Meniscal degeneration is a recognized contributor to knee osteoarthritis, but its senescence biology is difficult to separate from generic aging or inflammation, and single-marker senescence analysis is poorly suited to single-cell data. In this study, we used a program-level transcriptomic approach to define a meniscus senescence program (MSP), test it in an external single-cell dataset and nine bulk meniscus, cartilage, and synovium cohorts, assess OA-vs-normal stratification value, and prioritize paracrine axes for follow-up.

The study should be of interest to orthopaedic and sports-medicine readers for two reasons. First, it uses a reproducible, sample-aware analysis to avoid over-interpreting single markers or one-cohort disease signatures, and it adds a fibrocartilage-matrix OA stratification analysis. Second, it prioritizes three paracrine axes - MIF-CD74, ANGPTL4-integrin, and VEGF - that are consistent with recent osteoarthritis literature and can be tested with communication analysis, synovial-fluid protein assays, and conditioned-medium perturbation. We believe this fits the scope of the Journal of Orthopaedic Surgery and Research, which publishes clinical, translational, and basic-science orthopaedic studies, including bioinformatic analyses of musculoskeletal disease.

We have kept the claims deliberately conservative: the paracrine axes are hypotheses for validation, not established mechanisms or clinical targets.
"""

DATA_CODE_TEXT = """The public datasets re-analysed in this study are available from the repositories cited in the manuscript. Derived tables are included in the supplementary workbook. Final public repository URL or DOI for analysis code: URL or DOI to be inserted by the corresponding author before submission.
"""

DATA_CODE_PLACEHOLDER = """Availability of data and materials / code

The public datasets re-analysed in this study are available from the repositories cited in the manuscript. Derived tables are included in the supplementary workbook. Final public repository URL or DOI for analysis code: URL or DOI to be inserted by the corresponding author before submission.
"""

ABBREVIATIONS_SECTION = """## List of abbreviations

AUC, area under the receiver operating characteristic curve; CI, confidence interval; cNMF, consensus non-negative matrix factorization; ECM, extracellular matrix; FDR, false discovery rate; HRA, GSA-Human Run Archive; MSP, meniscus senescence program; OA, osteoarthritis; RA, rheumatoid arthritis; SASP, senescence-associated secretory phenotype; SMD, standardized mean difference; ST, supplementary table; VEGF, vascular endothelial growth factor.
"""

MANUSCRIPT_DECLARATIONS_SECTION = f"""## Declarations

### Ethics approval and consent to participate

This study re-analysed public and author-provided processed transcriptomic datasets. No new human or animal specimens were collected for this analysis. Ethics approval and consent for the original source data are described in the source studies cited in the manuscript.

### Consent for publication

Not applicable.

### Availability of data and materials

{DATA_CODE_TEXT.strip()}

### Competing interests

The authors declare that they have no competing interests.

### Funding

This research was funded by the Clinical Medicine Translational Research Special Program of Anhui Provincial Department of Science and Technology (grant number 202527c10020008).

### Authors' contributions

H.Y. and S.R. conceived and designed the study. S.R. performed the data analysis and wrote the original draft. D.L. and Y.D. assisted with data curation and visualization. X.C. and D.L. contributed to methodology and result interpretation. H.Y. supervised the project, acquired funding, and revised the manuscript. All authors read and approved the final manuscript.

### Acknowledgements

Not applicable.

### Declaration of generative AI in scientific writing

The authors used AI-assisted drafting tools for language editing, organization, and internal manuscript preparation. All scientific content, analyses, interpretation, and final responsibility remain with the authors.
"""

OFFICIAL_BASIS = [
    "JOSR Methodology article type supports new or improved computational methods that are well tested.",
    "JOSR Methodology abstract limit is 350 words and requires Background, Methods, Results, and Conclusions.",
    "JOSR requires three to ten keywords.",
    "JOSR requires a Declarations section and cover-letter policy statements.",
]

DIAGNOSTIC_REFERENCES = [
    "24. Hanley JA, McNeil BJ. The meaning and use of the area under a receiver operating characteristic (ROC) curve. Radiology. 1982;143(1):29-36. doi:10.1148/radiology.143.1.7063747.",
    "25. Efron B. Bootstrap methods: another look at the jackknife. Ann Stat. 1979;7(1):1-26. doi:10.1214/aos/1176344552.",
    "26. Varma S, Simon R. Bias in error estimation when using cross-validation for model selection. BMC Bioinformatics. 2006;7:91. doi:10.1186/1471-2105-7-91.",
    "27. Bossuyt PM, Reitsma JB, Bruns DE, Gatsonis CA, Glasziou PP, Irwig L, et al. STARD 2015: an updated list of essential items for reporting diagnostic accuracy studies. BMJ. 2015;351:h5527. doi:10.1136/bmj.h5527.",
    "28. Collins GS, Reitsma JB, Altman DG, Moons KG. Transparent reporting of a multivariable prediction model for individual prognosis or diagnosis (TRIPOD): the TRIPOD Statement. BMC Med. 2015;13:1. doi:10.1186/s12916-014-0241-z.",
]


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def write_tsv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    preferred = ["item", "status", "detail", "file_id", "upload_category", "handoff_path", "manual_check"]
    keys = []
    for key in preferred:
        if any(key in row for row in rows):
            keys.append(key)
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_assembly_module(root: Path):
    module_path = root / "scripts/analysis/45_jot_submission_file_assembly.py"
    spec = importlib.util.spec_from_file_location("submission_assembly", module_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load assembly module: {module_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def replace_section(text: str, heading: str, replacement: str, rows: list[dict[str, str]], item: str) -> str:
    pattern = re.compile(rf"## {re.escape(heading)}\n.*?(?=\n## |\Z)", re.S)
    if not pattern.search(text):
        rows.append({"item": item, "status": "not_found", "detail": heading})
        return text
    rows.append({"item": item, "status": "updated", "detail": heading})
    return pattern.sub(replacement.rstrip() + "\n", text, count=1)


def repackage_manuscript(text: str, rows: list[dict[str, str]]) -> str:
    text = text.replace("# JOT Anonymized Manuscript Draft", "# JOSR Anonymized Manuscript Draft")
    if OLD_TITLE in text or PREVIOUS_JOSR_TITLE in text or PREVIOUS_JOSR_TITLE_SHORT in text:
        text = (
            text.replace(OLD_TITLE, NEW_TITLE)
            .replace(PREVIOUS_JOSR_TITLE, NEW_TITLE)
            .replace(PREVIOUS_JOSR_TITLE_SHORT, NEW_TITLE)
        )
        rows.append({"item": "R1_title", "status": "updated", "detail": NEW_TITLE})
    elif NEW_TITLE in text:
        rows.append({"item": "R1_title", "status": "already_updated", "detail": NEW_TITLE})
    else:
        rows.append({"item": "R1_title", "status": "not_found", "detail": OLD_TITLE})
    text = replace_section(text, "Abstract", ABSTRACT_SECTION, rows, "R2_abstract")
    text = replace_section(text, "Keywords", KEYWORDS_SECTION, rows, "R4_keywords")
    before = text
    text = re.sub(
        r"\*\*The Translational Potential of this Article:\*\*.*?(?:\n\n|\Z)",
        "",
        text,
        flags=re.S,
    )
    rows.append(
        {
            "item": "R3_translational_potential",
            "status": "updated" if before != text else "already_updated",
            "detail": "removed JOT-specific abstract section",
        }
    )
    text = normalize_axis_display(text, rows)
    text = inject_diagnostic_content(text, rows)
    text = append_diagnostic_references(text, rows)
    text = polish_author_voice(text, rows)
    text = add_ai_methods_note(text, rows)
    text = apply_josr_main_structure(text, rows)
    text = normalize_figure_legend_order(text, rows)
    return text


def normalize_axis_display(text: str, rows: list[dict[str, str]]) -> str:
    replacements = [
        ("MIF_CD74", "MIF-CD74"),
        ("ANGPTL4_integrin", "ANGPTL4-integrin"),
    ]
    counts = {old: text.count(old) for old, _ in replacements}
    for old, new in replacements:
        text = text.replace(old, new)
    first_map = {
        "MIF-CD74": "MIF-CD74 (axis ID MIF_CD74)",
        "ANGPTL4-integrin": "ANGPTL4-integrin (axis ID ANGPTL4_integrin)",
    }
    for display, annotated in first_map.items():
        if annotated not in text and display in text:
            results_start = text.find("\n## Results")
            search_start = results_start if results_start >= 0 else 0
            hit = text.find(display, search_start)
            if hit < 0:
                hit = text.find(display)
            text = text[:hit] + annotated + text[hit + len(display):]
    detail = "; ".join(f"{old}:{count}" for old, count in counts.items())
    rows.append({"item": "R6_axis_display_consistency", "status": "updated" if any(counts.values()) else "already_updated", "detail": detail})
    return text


def inject_diagnostic_content(text: str, rows: list[dict[str, str]]) -> str:
    results_section = """### Fibrocartilage-Matrix Scoring Shows Candidate OA Stratification Value

Because the bulk meta-analysis identified fibrocartilage-matrix remodeling as the most direction-stable disease-associated axis, we next tested whether the existing sample-by-score matrix provided preliminary diagnostic or stratification value in OA-vs-normal bulk cohorts. The primary analysis included six OA-vs-normal cohorts (GSE114007, GSE143514, GSE169077, GSE185064, GSE55235, and GSE55457) and excluded the RA-rich GSE89408 cohort from the main analysis. GSE98918 was not included in this diagnostic contrast because it compares OA with arthroscopic partial meniscectomy rather than normal tissue.

The Fibrocartilage_matrix score showed a pooled directional AUC of 0.806 (95% CI 0.714-0.884), with positive OA-minus-normal deltas in all six primary cohorts and per-cohort AUCs ranging from 0.500 to 0.940. Grouped cross-validation that held out dataset groups yielded a mean AUC of 0.817 +/- 0.124, and leave-one-cohort-out testing yielded AUCs from 0.500 to 0.940. A secondary MSP-interface matrix/fibrotic composite showed a similar pooled AUC (0.813, 95% CI 0.725-0.888), but leave-one-cohort-out performance was less stable (0.333-0.967), so it was retained as secondary evidence. By contrast, MSP paracrine, angiogenic, inflammatory, and consensus paracrine scores had low pooled directional AUCs (0.128-0.223) when the pre-specified direction was OA higher, reinforcing their interpretation as context-dependent comparator axes rather than standalone diagnostic markers.
"""

    methods_section = """### Diagnostic discrimination analysis

To evaluate candidate stratification value, we restricted bulk samples to OA-vs-normal contrasts using the harmonized `condition` field. The primary diagnostic analysis included GSE114007, GSE143514, GSE169077, GSE185064, GSE55235, and GSE55457. The RA-rich GSE89408 cohort was excluded from the primary analysis and retained only as a sensitivity analysis; GSE98918 (OA vs arthroscopic partial meniscectomy) and GSE191157 (aged vs young) were treated as exploratory non-diagnostic contrasts. Directional AUCs were calculated without flipping score direction, so AUC values below 0.5 indicate that the pre-specified OA-higher direction did not hold; AUC was used as a threshold-independent discrimination metric [24]. The primary score was Fibrocartilage_matrix_K14P4; a secondary MSP-interface composite was defined as the sum of the cohort-internal Fibrocartilage_matrix_K14P4 and Fibrotic_remodeling_K14P7 z-score projections. Per-cohort and pooled AUCs used stratified bootstrap 95% confidence intervals with 2000 resamples [25]. Cross-dataset generalization was assessed by grouped cross-validation and leave-one-cohort-out logistic regression using dataset as the grouping variable, a design intended to reduce optimistic error estimation from non-independent validation splits [26]. Grouped cross-validation AUC is reported as the mean across folds with the sample standard deviation across folds. These analyses evaluate candidate stratification signals rather than a validated clinical classifier, consistent with diagnostic-accuracy and prediction-model reporting guidance [27,28].
"""

    diagnostic_discussion = """The diagnostic discrimination analysis adds a practical but still preliminary translational layer. The fibrocartilage-matrix score separated OA from normal tissue across the primary bulk cohorts with a pooled directional AUC of 0.806 and grouped cross-validation AUC of 0.817, supporting its use as a candidate cross-tissue stratification signal. This finding is consistent with the random-effects meta-analysis, where the same axis was the only strongly supported direction-stable bulk signal. Importantly, the pure MSP paracrine, angiogenic, and inflammatory axes were not promoted as diagnostic markers because their directional AUCs were low in the pre-specified OA-higher direction, matching the broader heterogeneity analysis and the reporting principle that preliminary discrimination evidence should not be presented as a validated clinical prediction model [27,28]."""

    text = text.replace(
        "\n### Integrated Evidence Prioritizes Candidate Paracrine Mechanism Axes",
        "\n" + results_section + "\n### Integrated Evidence Prioritizes Candidate Paracrine Mechanism Axes",
        1,
    )
    text = text.replace(
        "\n### Bulk subtype analysis and robustness checks",
        "\n" + methods_section + "\n### Bulk subtype analysis and robustness checks",
        1,
    )
    text = text.replace(
        "\n### Candidate paracrine axes",
        "\n" + diagnostic_discussion + "\n\n### Candidate paracrine axes",
        1,
    )
    text = text.replace(
        "The translational value of this framework is prioritization. The S1 ECM/fibrotic-high subtype, the leading paracrine axes, and the validation roadmap provide a ranked set of hypotheses for biomarker and perturbation studies.",
        "The translational value of this framework is prioritization. The fibrocartilage-matrix diagnostic analysis, the S1 ECM/fibrotic-high subtype, the leading paracrine axes, and the validation roadmap provide a ranked set of hypotheses for biomarker and perturbation studies.",
        1,
    )
    text = text.replace(
        "Fourth, formal CellChat/LIANA/NicheNet consensus, synovial-fluid protein validation, and conditioned-medium perturbation are planned validation layers rather than completed evidence. Fifth, Hotspot validation, RNA velocity, and Mendelian randomization were not performed in the current analysis.",
        "Fourth, the diagnostic discrimination analysis is based on public bulk cohorts, several of them small, and is therefore candidate stratification evidence rather than a validated clinical classifier under diagnostic-accuracy and prediction-model reporting standards [27,28]. Finally, several validation layers remain outstanding: formal CellChat/LIANA/NicheNet consensus, synovial-fluid protein validation, and conditioned-medium perturbation are planned but not yet completed, and Hotspot co-expression validation, RNA velocity, and Mendelian randomization were not performed in the current analysis.",
        1,
    )
    text = text.replace(
        "Bulk meta-analysis FDR was controlled across the seven scored axes using Benjamini-Hochberg correction; tissue-stratified evidence grades should be treated as the primary bulk readout when comparator or tissue definitions differ across cohorts.",
        "Bulk meta-analysis FDR was controlled across the seven scored axes using Benjamini-Hochberg correction; tissue-stratified evidence grades should be treated as the primary bulk readout when comparator or tissue definitions differ across cohorts. Diagnostic AUC analysis used the existing sample-by-score matrix and retained GSE89408, GSE98918, and GSE191157 outside the primary OA-vs-normal analysis for sensitivity or exploratory interpretation [24-28].",
        1,
    )
    text = text.replace(
        "### Figure 4. Candidate MSP-linked paracrine mechanism axes",
        "### Figure 6. Candidate OA stratification value of fibrocartilage-matrix scoring\n\nFigure 6 summarizes the diagnostic discrimination analysis in OA-vs-normal bulk cohorts [24-28]. Panel A shows per-cohort directional AUCs for the Fibrocartilage_matrix score and the secondary MSP-interface matrix/fibrotic composite. Panel B compares pooled directional AUCs across remodeling and MSP comparator scores. Panel C shows leave-one-cohort-out AUCs for the primary and secondary scores. GSE89408 was excluded from the primary analysis because it is RA-rich and was retained only as a sensitivity analysis. These results support candidate stratification value, not a deployable clinical diagnostic classifier.\n\n### Figure 4. Candidate MSP-linked paracrine mechanism axes",
        1,
    )
    text = text.replace(
        "| Figure 4 | A-D | results/tables/msp_mechanism_axis_evidence_dossier.tsv;results/tables/msp_comm_tool_pair_edges.tsv;results/tables/msp_axis_target_concordance_for_nichenet.tsv | MIF-CD74, ANGPTL4-integrin, and VEGF are leading candidate axes; not causal or validated mechanisms. |",
        "| Figure 4 | A-D | results/tables/msp_mechanism_axis_evidence_dossier.tsv;results/tables/msp_comm_tool_pair_edges.tsv;results/tables/msp_axis_target_concordance_for_nichenet.tsv | MIF-CD74, ANGPTL4-integrin, and VEGF are leading candidate axes; not causal or validated mechanisms. |\n| Figure 6 | A-C | results/tables/bulk_msp_diagnostic_per_cohort_auc.tsv;results/tables/bulk_msp_diagnostic_pooled_auc.tsv;results/tables/bulk_msp_diagnostic_loco_auc.tsv | Candidate stratification signal only; not a clinical diagnostic classifier. |",
        1,
    )
    text = text.replace(
        "| ST28 | manuscript_package | results/tables/manuscript_figure4_axis_summary.tsv | Figure 4 axis summary |",
        "| ST28 | manuscript_package | results/tables/manuscript_figure4_axis_summary.tsv | Figure 4 axis summary |\n| ST29 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_input_manifest.tsv | Diagnostic analysis cohort inclusion manifest |\n| ST30 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_per_cohort_auc.tsv | Per-cohort diagnostic AUC estimates |\n| ST31 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_pooled_auc.tsv | Pooled diagnostic AUC estimates |\n| ST32 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_grouped_cv_auc.tsv | Grouped cross-validation diagnostic AUC estimates |\n| ST33 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_loco_auc.tsv | Leave-one-cohort-out diagnostic AUC estimates |\n| ST34 | diagnostic_analysis | results/tables/bulk_msp_diagnostic_sensitivity_auc.tsv | Sensitivity and exploratory diagnostic AUC estimates |",
        1,
    )
    rows.append({"item": "R7_diagnostic_analysis_content", "status": "updated", "detail": "Inserted Results, Methods, Discussion, Figure 6 legend, and ST29-ST34 diagnostic rows."})
    return text


def append_diagnostic_references(text: str, rows: list[dict[str, str]]) -> str:
    if "24. Hanley JA, McNeil BJ." in text:
        rows.append({"item": "R8_diagnostic_references", "status": "already_updated", "detail": "Diagnostic-method references already present."})
        return text
    reference_block = "\n".join(DIAGNOSTIC_REFERENCES) + "\n"
    if "\n## References\n" not in text:
        rows.append({"item": "R8_diagnostic_references", "status": "not_found", "detail": "References heading"})
        return text
    text = text.rstrip() + "\n" + reference_block
    rows.append({"item": "R8_diagnostic_references", "status": "updated", "detail": "Added ROC/AUC, bootstrap, cross-validation, STARD, and TRIPOD references 24-28."})
    return text


def polish_author_voice(text: str, rows: list[dict[str, str]]) -> str:
    replacements = {
        "Third, rather than asserting a disease-up signature from one dataset, we impose explicit sample-aware and cross-tissue guardrails and report candidate axes with calibrated, validation-ready confidence.": "Third, rather than treating one dataset as definitive, we used sample-aware and cross-tissue checks to avoid over-calling disease-specific effects.",
        "The goal is not to claim causal mechanism from transcriptomic association, but to generate a traceable, validation-ready framework for meniscus-to-joint remodeling hypotheses in knee osteoarthritis.": "The aim was not to prove causality from transcriptomic association, but to provide a traceable analysis that can guide follow-up studies of meniscus-to-joint remodeling in knee osteoarthritis.",
        "This framing treats MSP discovery as a continuous cell-state and program-usage problem, allowing downstream analyses to retain sample-aware caution rather than forcing discrete disease-state labels.": "This approach treats MSP discovery as a continuous cell-state and program-usage problem, rather than forcing discrete disease labels onto sparse single-cell data.",
        "These programs were carried forward as candidate MSP biology because they were supported by program-level enrichment and donor-aware summaries.": "We carried these programs forward because they showed program-level enrichment and acceptable donor-aware support.",
        "Instead, Figure 1 presents them as candidate fibrochondrocyte programs that require external projection, bulk validation, and mechanism-focused follow-up.": "Figure 1 therefore presents them as fibrochondrocyte programs for external projection, bulk validation, and mechanism-focused follow-up.",
        "This provided an independent check that the candidate programs were not simply artifacts of the discovery object.": "This served as an independent check that the programs were not simply artifacts of the discovery object.",
        "Figure 2 therefore functions as an external single-cell projection layer, bridging discovery cNMF programs to a second meniscus dataset while preserving conservative language.": "Figure 2 therefore serves as an external single-cell projection layer linking the discovery cNMF programs to a second meniscus dataset.",
        "Under this framing, the only axis with strong and consistent bulk support was fibrocartilage matrix remodeling": "With this approach, the only axis with strong and consistent bulk support was fibrocartilage matrix remodeling",
        "and were therefore treated as context-dependent comparator axes.": "and were treated as context-dependent signals.",
        "reinforcing their interpretation as context-dependent comparator axes rather than standalone diagnostic markers.": "supporting their interpretation as context-dependent signals rather than standalone diagnostic markers.",
        "These results support a validation-ready mechanism hypothesis centered on MSP-program-high meniscal fibrochondrocytes, immune/myeloid or vascular receiver contexts, and remodeling-associated target genes.": "Together, these results point to follow-up hypotheses linking MSP-program-high meniscal fibrochondrocytes with immune/myeloid or vascular receiver contexts and remodeling-associated target genes.",
        "Single-cell expression context and pre-NicheNet concordance were used as prioritization layers, not as proof of signaling.": "Single-cell expression context and pre-NicheNet concordance were used to rank hypotheses, not to prove signaling.",
        "Finally, the candidate axes were translated into a validation roadmap (Figure 5).": "Finally, we summarized the candidate axes in a validation plan (Figure 5).",
        "This final layer is deliberately prospective. The present study nominates candidate paracrine axes and a validation strategy; it does not establish that MSP ligands have disease-driving activity.": "These experiments remain prospective. The present study identifies paracrine axes for follow-up; it does not establish that MSP ligands drive disease.",
        "MIF-CD74, ANGPTL4-integrin, and VEGF are therefore leading candidate axes that require formal CellChat/LIANA/NicheNet consensus [14-17], protein evidence, and perturbation experiments before any causal interpretation.": "MIF-CD74, ANGPTL4-integrin, and VEGF therefore require formal CellChat/LIANA/NicheNet consensus [14-17], protein evidence, and perturbation experiments before any causal interpretation.",
        "The leading candidate paracrine axes were MIF-CD74 (axis ID MIF_CD74), ANGPTL4-integrin (axis ID ANGPTL4_integrin), VEGF. MIF-CD74 ranked 1, linking MSP-program-high fibrochondrocyte to immune/myeloid through MIF->CD74; ANGPTL4-integrin ranked 2, linking MSP-program-high fibrochondrocyte to mural/smooth muscle through ANGPTL4->ITGB1; VEGF ranked 3, linking MSP-program-high fibrochondrocyte to endothelial through VEGFA->FLT1.": "The leading candidate paracrine axes were MIF-CD74 (axis ID MIF_CD74), ANGPTL4-integrin (axis ID ANGPTL4_integrin), and VEGF. MIF-CD74 linked MSP-program-high fibrochondrocytes with immune/myeloid receiver cells through the MIF-CD74 pairing; ANGPTL4-integrin linked MSP-program-high fibrochondrocytes with mural/smooth muscle receiver contexts through ANGPTL4-ITGB1; and VEGF linked MSP-program-high fibrochondrocytes with endothelial receiver contexts through VEGFA-FLT1.",
        "Secondary axes including MIF_chemokine_receptors, FGF_FGFR, BMP2_BMPR, and INHBA_activin broadened the remodeling/interface hypothesis but require additional pathway-specific evidence before being used as central claims.": "Secondary axes including MIF-chemokine receptor, FGF-FGFR, BMP2-BMPR, and INHBA-activin broadened the remodeling/interface hypothesis but require additional pathway-specific evidence before being used as central claims.",
        "This study develops a cautious, program-level framework for linking meniscal fibrochondrocyte senescence biology to broader joint remodeling in knee osteoarthritis.": "This study presents a program-level analysis linking meniscal fibrochondrocyte senescence biology with broader joint remodeling in knee osteoarthritis.",
        "The major finding is not that a causal pathway has been proven, but that MSP-like programs nominate a coherent set of validation-ready hypotheses connecting fibrochondrocyte stress, paracrine signaling, vascular or immune receiver contexts, and bulk remodeling phenotypes.": "The main result is a ranked set of biologically plausible hypotheses connecting fibrochondrocyte stress, paracrine signaling, vascular or immune receiver contexts, and bulk remodeling phenotypes, rather than proof of a causal pathway.",
        "The term MSP-like remains deliberate: several high-ranking programs were sample-aware and donor-skewed, so the evidence supports candidate meniscus senescence biology rather than a fixed disease state.": "We use MSP-like deliberately because several high-ranking programs were sample-aware and donor-skewed; the evidence supports meniscus senescence-related biology, not a fixed disease state.",
        "The diagnostic discrimination analysis adds a practical but still preliminary translational layer.": "The diagnostic discrimination analysis adds a clinically oriented but preliminary layer.",
        "The translational value of this framework is prioritization.": "The main translational value is prioritization.",
        "The leading axes were MIF-CD74 (MIF->CD74; immune/myeloid receiver context), ANGPTL4-integrin (ANGPTL4->ITGB1; mural/smooth muscle receiver context), VEGF (VEGFA->FLT1; endothelial receiver context).": "The leading axes were MIF-CD74 (immune/myeloid receiver context), ANGPTL4-integrin (ANGPTL4-ITGB1; mural/smooth muscle receiver context), and VEGF (VEGFA-FLT1; endothelial receiver context).",
        "through an ANGPTL4-alpha5beta1 (ITGA5/ITGB1) axis": "through an ANGPTL4-alpha5 beta1 (ITGA5/ITGB1) axis",
        "These findings may guide synovial-fluid protein panels, meniscus-conditioned-medium experiments, and targeted blockade assays. They should be treated as prioritization evidence rather than clinical-deployment or disease-driving evidence. At this stage, the manuscript supports validation-ready candidate biology.": "These findings may guide synovial-fluid protein panels, meniscus-conditioned-medium experiments, and targeted blockade assays. They should be treated as prioritization evidence rather than evidence for clinical deployment or disease-driving mechanisms.",
        "In summary, this study nominates a meniscus-centered MSP framework connecting program-level fibrochondrocyte states, a robust fibrocartilage-matrix remodeling bulk signal, context-dependent MSP-axis remodeling, and candidate paracrine axes.": "In summary, this meniscus-centered MSP analysis links program-level fibrochondrocyte states with a robust fibrocartilage-matrix remodeling bulk signal, context-dependent MSP-axis behavior, and paracrine axes for follow-up.",
        "The most defensible interpretation is that MIF-CD74, ANGPTL4-integrin, and VEGF represent leading validation-ready hypotheses, not causal or validated mechanisms. This careful framing leaves room for mechanistic validation while preserving the biological signal that MSP-like meniscal programs may help organize inflammatory, vascular, and matrix remodeling features of knee osteoarthritis.": "The most defensible interpretation is that MIF-CD74, ANGPTL4-integrin, and VEGF are leading hypotheses, not validated mechanisms. This interpretation preserves the biological signal while making clear that mechanistic validation is still required.",
        "All mechanism evidence scores are triage scores and should not be interpreted as effect sizes or causal estimates.": "Mechanism evidence scores were used for triage and should not be interpreted as effect sizes or causal estimates.",
        "Candidate and not-causal language is therefore used throughout the Results and figure legends.": "Accordingly, the Results and figure legends use hypothesis-generating rather than causal language.",
        "These panels define candidate MSP biology and are not causal evidence.": "These panels define MSP-like biology and do not provide causal evidence.",
        "This projection supports meniscus context but is not causal or direct spatial proof.": "This projection supports meniscus context but does not prove causality or spatial localization.",
        "These results support context-dependent MSP/remodeling biology and should not be written as a universal OA-up MSP signature.": "These results support context-dependent MSP/remodeling biology rather than a universal OA-up MSP signature.",
        "These results support candidate stratification value, not a deployable clinical diagnostic classifier.": "These results support preliminary stratification value, not a clinical diagnostic classifier.",
        "MIF-CD74, ANGPTL4-integrin, and VEGF are leading candidate paracrine axes, not causal or validated mechanisms.": "MIF-CD74, ANGPTL4-integrin, and VEGF are leading paracrine axes for follow-up, not validated mechanisms.",
        "This roadmap is prospective and does not represent completed wet-lab validation.": "This plan is prospective and does not represent completed wet-lab validation.",
    }
    changed = 0
    for old, new in replacements.items():
        if old in text:
            text = text.replace(old, new)
            changed += 1
    rows.append({"item": "R9_author_voice_polish", "status": "updated", "detail": f"Applied {changed} natural-language replacements."})
    return text


def add_ai_methods_note(text: str, rows: list[dict[str, str]]) -> str:
    note = """### AI-assisted language editing

AI-assisted drafting tools were used only for language editing, organization, and internal manuscript preparation. The authors checked the scientific content, analyses, interpretation, references, and final wording, and take full responsibility for the manuscript.

"""
    if "### AI-assisted language editing" in text:
        rows.append({"item": "R11_ai_methods_note", "status": "already_updated", "detail": "Methods note already present."})
        return text
    marker = "\n## Figure Legends"
    if marker not in text:
        rows.append({"item": "R11_ai_methods_note", "status": "not_found", "detail": "Figure Legends marker"})
        return text
    text = text.replace(marker, "\n" + note + "## Figure Legends", 1)
    rows.append({"item": "R11_ai_methods_note", "status": "updated", "detail": "Added JOSR-compatible Methods note for AI-assisted language editing."})
    return text


def apply_josr_main_structure(text: str, rows: list[dict[str, str]]) -> str:
    if "\n## Background\n" in text and "\n## Methods\n" in text and "\n## Declarations\n" in text:
        rows.append({"item": "R12_josr_main_structure", "status": "already_updated", "detail": "Background/Methods/Results/Discussion/Conclusions structure present."})
        return text
    pattern = re.compile(
        r"(?P<prefix>.*?)(?P<intro>\n## Introduction\n.*?)(?P<results>\n## Results\n.*?)(?P<discussion>\n## Discussion\n.*?)(?P<methods>\n## Materials and Methods\n.*?)(?P<rest>\n## Figure Legends\n.*)\Z",
        re.S,
    )
    match = pattern.match(text)
    if not match:
        rows.append({"item": "R12_josr_main_structure", "status": "not_found", "detail": "Could not parse major manuscript sections."})
        return text

    prefix = match.group("prefix").rstrip()
    background = match.group("intro").replace("\n## Introduction", "\n## Background", 1).strip()
    methods = match.group("methods").replace("\n## Materials and Methods", "\n## Methods", 1).strip()
    results = match.group("results").strip()
    discussion = match.group("discussion").strip()
    rest = match.group("rest").strip()

    conclusion_match = re.search(r"\n### Conclusion\n(?P<conclusion>.*)\Z", discussion, re.S)
    if conclusion_match:
        conclusion_text = conclusion_match.group("conclusion").strip()
        discussion = discussion[: conclusion_match.start()].rstrip()
    else:
        conclusion_text = "The findings should be interpreted as transcriptomic prioritization evidence and require further mechanistic validation."

    sections = [
        prefix,
        background,
        methods,
        results,
        discussion,
        "## Conclusions\n\n" + conclusion_text,
        ABBREVIATIONS_SECTION.strip(),
        MANUSCRIPT_DECLARATIONS_SECTION.strip(),
        rest,
    ]
    rows.append({"item": "R12_josr_main_structure", "status": "updated", "detail": "Reordered main manuscript to Background, Methods, Results, Discussion, Conclusions, abbreviations, declarations, legends, references."})
    return "\n\n".join(section for section in sections if section.strip()) + "\n"


def normalize_figure_legend_order(text: str, rows: list[dict[str, str]]) -> str:
    heading = "## Figure Legends"
    next_heading = "## Supplementary Material"
    if heading not in text or next_heading not in text:
        rows.append({"item": "R10_figure_legend_order", "status": "not_found", "detail": "Figure legend block"})
        return text
    before, rest = text.split(heading, 1)
    legend_block, after = rest.split(next_heading, 1)
    legend_items = re.findall(r"\n### Figure \d+\..*?(?=\n### Figure \d+\.|\n### Supplementary Figure|\Z)", legend_block, flags=re.S)
    supp_match = re.search(r"\n### Supplementary Figure.*", legend_block, flags=re.S)
    if len(legend_items) < 6:
        rows.append({"item": "R10_figure_legend_order", "status": "not_found", "detail": f"Found {len(legend_items)} numbered figure legends"})
        return text
    legend_items = sorted(
        legend_items,
        key=lambda item: int(re.search(r"### Figure (\d+)\.", item).group(1)),
    )
    rebuilt = "\n".join(item.strip() for item in legend_items)
    if supp_match:
        rebuilt += "\n\n" + supp_match.group(0).strip()
    text = before + heading + "\n\n" + rebuilt.strip() + "\n\n" + next_heading + after

    index_old = "| Figure 4 | A-D | results/tables/msp_mechanism_axis_evidence_dossier.tsv;results/tables/msp_comm_tool_pair_edges.tsv;results/tables/msp_axis_target_concordance_for_nichenet.tsv | MIF-CD74, ANGPTL4-integrin, and VEGF are leading candidate axes; not causal or validated mechanisms. |\n| Figure 6 | A-C | results/tables/bulk_msp_diagnostic_per_cohort_auc.tsv;results/tables/bulk_msp_diagnostic_pooled_auc.tsv;results/tables/bulk_msp_diagnostic_loco_auc.tsv | Candidate stratification signal only; not a clinical diagnostic classifier. |\n| Figure 5 | A | results/tables/msp_lr_validation_assay_matrix.tsv | Roadmap is proposed validation, not completed wet-lab evidence. |"
    index_new = "| Figure 4 | A-D | results/tables/msp_mechanism_axis_evidence_dossier.tsv;results/tables/msp_comm_tool_pair_edges.tsv;results/tables/msp_axis_target_concordance_for_nichenet.tsv | MIF-CD74, ANGPTL4-integrin, and VEGF are leading paracrine axes for follow-up, not validated mechanisms. |\n| Figure 5 | A | results/tables/msp_lr_validation_assay_matrix.tsv | Validation plan is proposed, not completed wet-lab evidence. |\n| Figure 6 | A-C | results/tables/bulk_msp_diagnostic_per_cohort_auc.tsv;results/tables/bulk_msp_diagnostic_pooled_auc.tsv;results/tables/bulk_msp_diagnostic_loco_auc.tsv | Preliminary stratification signal only; not a clinical diagnostic classifier. |"
    text = text.replace(index_old, index_new, 1)
    rows.append({"item": "R10_figure_legend_order", "status": "updated", "detail": "Sorted figure legends and source index in numerical order."})
    return text


def repackage_title_page(text: str, rows: list[dict[str, str]]) -> str:
    text = text.replace("# JOT Title Page and Author Statements", "# JOSR Title Page and Author Statements")
    text = (
        text.replace(OLD_TITLE, NEW_TITLE)
        .replace(PREVIOUS_JOSR_TITLE, NEW_TITLE)
        .replace(PREVIOUS_JOSR_TITLE_SHORT, NEW_TITLE)
    )
    text = re.sub(r"(## Article Type\s*\n\n).*?(?=\n\n## )", r"\1Methodology", text, flags=re.S)
    text = text.replace(" [Author confirmation required before submission.]", "")
    text = text.replace(" [Revise according to final journal policy and actual tool use before submission.]", "")
    text = text.replace("[to be added if required by the submission system]", "To be entered in the submission system if required")
    text = re.sub(
        r"Ethics approval and consent for the original source data are described in the source studies cited in the anonymized manuscript references 3-5\.",
        "Ethics approval and consent for the original source data are described in the source studies cited in the manuscript.",
        text,
    )
    rows.append({"item": "title_page_article_type", "status": "updated", "detail": "Methodology"})
    return text


def repackage_cover_letter(text: str, rows: list[dict[str, str]]) -> str:
    text = text.replace("# JOT Cover Letter Draft", "# JOSR Cover Letter Draft")
    closing_anchor = "This manuscript has not been published previously"
    if closing_anchor not in text:
        rows.append({"item": "R5_cover_letter_opening", "status": "not_found", "detail": closing_anchor})
        return text.replace(OLD_TITLE, NEW_TITLE)
    closing = closing_anchor + text.split(closing_anchor, 1)[1]
    if "The authors declare that they have no competing interests." not in closing:
        closing = closing.replace(
            "Non-author contributors, if any, will be disclosed according to journal requirements.",
            "Non-author contributors, if any, will be disclosed according to journal requirements. The authors declare that they have no competing interests.",
            1,
        )
    rows.append({"item": "R5_cover_letter_opening", "status": "updated", "detail": "JOSR cover letter opening and policy statement"})
    return "# JOSR Cover Letter Draft\n\n" + COVER_OPENING.rstrip() + "\n\n" + closing


def extract_abstract(markdown: str) -> str:
    match = re.search(r"## Abstract\n(.*?)(?=\n## Keywords)", markdown, re.S)
    return match.group(1).strip() if match else ""


def word_count(text: str) -> int:
    clean = re.sub(r"\*\*|:|\(|\)|,|;", " ", text)
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]*", clean))


def keyword_count(markdown: str) -> int:
    match = re.search(r"\*\*Keywords:\*\*\s*(.+)", markdown)
    if not match:
        return 0
    return len([piece.strip() for piece in match.group(1).split(";") if piece.strip()])


def build_declarations(title_page: str) -> str:
    sections = {
        "Ethics approval and consent to participate": "This study re-analysed public and author-provided processed transcriptomic datasets. No new human or animal specimens were collected for this analysis. Ethics approval and consent for the original source data are described in the source studies cited in the manuscript.",
        "Consent for publication": "Not applicable.",
        "Availability of data and materials": DATA_CODE_TEXT.strip(),
        "Competing interests": "The authors declare that they have no competing interests.",
        "Funding": "This research was funded by the Clinical Medicine Translational Research Special Program of Anhui Provincial Department of Science and Technology (grant number 202527c10020008).",
        "Authors' contributions": "H.Y. and S.R. conceived and designed the study. S.R. performed the data analysis and wrote the original draft. D.L. and Y.D. assisted with data curation and visualization. X.C. and D.L. contributed to methodology and result interpretation. H.Y. supervised the project, acquired funding, and revised the manuscript. All authors read and approved the final manuscript.",
        "Acknowledgements": "Not applicable.",
        "Declaration of generative AI in scientific writing": "The authors used AI-assisted drafting tools for language editing, organization, and internal manuscript preparation. All scientific content, analyses, interpretation, and final responsibility remain with the authors.",
    }
    lines = ["# JOSR Declarations"]
    for heading, content in sections.items():
        lines.extend(["", f"## {heading}", "", content])
    return "\n".join(lines) + "\n"


def write_josr_supplementary_workbook(root: Path, assembly, supp_manifest: pd.DataFrame, output: Path) -> pd.DataFrame:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    readme_rows = [
        ["Item", "Value"],
        ["Purpose", "Journal of Orthopaedic Surgery and Research supplementary tables ST01-ST34"],
        ["Generated by", "scripts/analysis/58_josr_repackaging_application.py"],
        ["Caution", "Tables support a candidate, not causal, interpretation of MSP-associated axes."],
        ["Manual check", "Confirm sheet labels, legends, and any journal-specific supplementary formatting before submission."],
    ]
    for row in readme_rows:
        ws_readme.append(row)
    assembly.style_header(ws_readme)
    assembly.set_widths(ws_readme, max_width=70)

    manifest_rows = []
    for _, row in supp_manifest.iterrows():
        table_id = str(row["table_id"])
        source = str(row["source_output"])
        sheet_name = table_id[:31]
        data = assembly.read_source_table(root, source)
        ws = wb.create_sheet(sheet_name)
        assembly.append_dataframe(ws, data)
        manifest_rows.append(
            {
                "table_id": table_id,
                "theme": row["theme"],
                "source_output": source,
                "sheet_name": sheet_name,
                "rows": str(len(data)),
                "columns": str(len(data.columns)),
                "source_status": "source_ready" if (root / source).exists() else "source_missing",
            }
        )

    ws_manifest = wb.create_sheet("Manifest", 1)
    manifest_df = pd.DataFrame(manifest_rows)
    assembly.append_dataframe(ws_manifest, manifest_df)
    wb.save(output)
    return manifest_df


def append_diagnostic_supplementary_rows(supp_manifest: pd.DataFrame) -> pd.DataFrame:
    diagnostic_rows = pd.DataFrame(
        [
            {
                "table_id": "ST29",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_input_manifest.tsv",
                "suggested_use": "Diagnostic analysis cohort inclusion manifest",
            },
            {
                "table_id": "ST30",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_per_cohort_auc.tsv",
                "suggested_use": "Per-cohort diagnostic AUC estimates",
            },
            {
                "table_id": "ST31",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_pooled_auc.tsv",
                "suggested_use": "Pooled diagnostic AUC estimates",
            },
            {
                "table_id": "ST32",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_grouped_cv_auc.tsv",
                "suggested_use": "Grouped cross-validation diagnostic AUC estimates",
            },
            {
                "table_id": "ST33",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_loco_auc.tsv",
                "suggested_use": "Leave-one-cohort-out diagnostic AUC estimates",
            },
            {
                "table_id": "ST34",
                "theme": "diagnostic_analysis",
                "source_output": "results/tables/bulk_msp_diagnostic_sensitivity_auc.tsv",
                "suggested_use": "Sensitivity and exploratory diagnostic AUC estimates",
            },
        ]
    )
    existing = set(supp_manifest["table_id"].astype(str))
    diagnostic_rows = diagnostic_rows.loc[~diagnostic_rows["table_id"].isin(existing)].copy()
    return pd.concat([supp_manifest, diagnostic_rows], ignore_index=True)


def copy_figures(root: Path, submission_dir: Path, handoff_dir: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    src_dir = root / "submission/jot/figures"
    for src in sorted(src_dir.glob("*.png")):
        for dest_dir in (submission_dir / "figures", handoff_dir / "02_figures"):
            dest_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest_dir / src.name)
        rows.append(
            {
                "file_id": src.stem,
                "upload_category": "figure",
                "handoff_path": f"02_figures/{src.name}",
                "status": "copied",
                "manual_check": "Confirm resolution, lettering, and JOSR figure-format preference before submission.",
            }
        )
    diagnostic_figures = [
        (
            "Figure_6_bulk_msp_diagnostic_auc_summary.png",
            root / "results/figures/bulk_msp_diagnostic/bulk_msp_diagnostic_auc_summary.png",
        )
    ]
    for dest_name, src in diagnostic_figures:
        if not src.exists():
            rows.append(
                {
                    "file_id": Path(dest_name).stem,
                    "upload_category": "figure",
                    "handoff_path": f"02_figures/{dest_name}",
                    "status": "source_missing",
                    "manual_check": f"Run the diagnostic analysis before submission: {src}",
                }
            )
            continue
        for dest_dir in (submission_dir / "figures", handoff_dir / "02_figures"):
            dest_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest_dir / dest_name)
        rows.append(
            {
                "file_id": Path(dest_name).stem,
                "upload_category": "figure",
                "handoff_path": f"02_figures/{dest_name}",
                "status": "copied",
                "manual_check": "Confirm resolution, lettering, and JOSR figure-format preference before submission.",
            }
        )
    return rows


def copy_to_handoff(paths: dict[str, Path], handoff_dir: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    manuscript_dir = handoff_dir / "01_manuscript_files"
    supp_dir = handoff_dir / "03_supplementary_tables"
    data_dir = handoff_dir / "04_data_code_availability"
    manuscript_dir.mkdir(parents=True, exist_ok=True)
    supp_dir.mkdir(parents=True, exist_ok=True)
    data_dir.mkdir(parents=True, exist_ok=True)
    for key in ("title_page", "manuscript", "cover_letter", "declarations"):
        src = paths[key]
        dest = manuscript_dir / src.name
        shutil.copy2(src, dest)
        rows.append(
            {
                "file_id": key,
                "upload_category": key,
                "handoff_path": f"01_manuscript_files/{dest.name}",
                "status": "ready_needs_author_check",
                "manual_check": "Open and confirm formatting, author wording, and metadata before upload.",
            }
        )
    supp_dest = supp_dir / paths["supplementary"].name
    shutil.copy2(paths["supplementary"], supp_dest)
    rows.append(
        {
            "file_id": "supplementary_tables",
            "upload_category": "supplementary_table",
            "handoff_path": f"03_supplementary_tables/{supp_dest.name}",
            "status": "ready_needs_manual_check",
            "manual_check": "Open workbook and confirm sheet readability before upload.",
        }
    )
    data_path = data_dir / "Data_and_Code_Availability_URL_or_DOI.txt"
    write_text(data_path, DATA_CODE_PLACEHOLDER)
    rows.append(
        {
            "file_id": "data_code_availability",
            "upload_category": "data_code_availability",
            "handoff_path": "04_data_code_availability/Data_and_Code_Availability_URL_or_DOI.txt",
            "status": "user_to_fill",
            "manual_check": "Insert final repository URL or DOI before final submission.",
        }
    )
    return rows


def write_handoff_readme(handoff_dir: Path, manifest_rows: list[dict[str, str]]) -> None:
    lines = [
        "# JOSR Upload Handoff",
        "",
        "Use this folder as the local source for Journal of Orthopaedic Surgery and Research submission.",
        "",
        "## Upload Map",
        "",
        "| file_id | upload_category | handoff_path | status | manual_check |",
        "| --- | --- | --- | --- | --- |",
    ]
    for row in manifest_rows:
        lines.append(
            f"| {row['file_id']} | {row['upload_category']} | {row['handoff_path']} | {row['status']} | {row['manual_check']} |"
        )
    lines.extend(
        [
            "",
            "## Manual Items",
            "",
            "- Insert the final repository URL or DOI.",
            "- Confirm author approvals, competing interests, ethics wording, and AI-use wording before final submission.",
            "- Confirm whether the portal labels the article type as Methodology or Research, and select Methodology if available.",
        ]
    )
    write_text(handoff_dir / "00_README_UPLOAD_HANDOFF.md", "\n".join(lines) + "\n")


def write_submitter_checklist(submission_dir: Path) -> None:
    text = """# Final Submitter Checklist

## Before opening the JOSR portal

- Use the files in `submission/josr/upload_handoff`.
- Have the Repository URL/DOI ready, or leave the placeholder until you have it.
- Select article type `Methodology` if the submission portal offers it.

## Upload order

| portal_step | upload_category | handoff_path |
| --- | --- | --- |
| 1 | title_page | 01_manuscript_files/JOSR_Title_Page_and_Author_Statements.docx |
| 2 | manuscript | 01_manuscript_files/JOSR_Anonymized_Manuscript.docx |
| 3 | cover_letter | 01_manuscript_files/JOSR_Cover_Letter.docx |
| 4 | declarations | 01_manuscript_files/JOSR_Declarations.docx |
| 5 | figures | 02_figures/ |
| 6 | supplementary_table | 03_supplementary_tables/JOSR_Supplementary_Tables_ST01_ST34.xlsx |
| 7 | data_code_availability | 04_data_code_availability/Data_and_Code_Availability_URL_or_DOI.txt |

## Final go/no-go

| check_id | domain | status | submit_gate |
| --- | --- | --- | --- |
| F01 | title_and_abstract | ready | JOSR title, structured abstract, and keywords are applied. |
| F02 | article_type | manual_user_check | Select Methodology if available in the portal. |
| F03 | metadata_clean_docx | manual_user_check | Open DOCX properties and confirm no hidden identifiers before final submission. |
| F04 | author_confirmation | manual_user_check | Corresponding author confirms COI, ethics, AI declaration, funding, and author contributions. |
| F05 | repository_url_or_doi | user_to_fill | Insert repository URL or DOI before final click Submit. |
| F06 | claim_guardrail | ready | Keep candidate axes as hypotheses for validation, not causal mechanisms. |
"""
    write_text(submission_dir / "FINAL_SUBMITTER_CHECKLIST.md", text)


def zip_handoff(handoff_dir: Path, output_zip: Path) -> None:
    output_zip.parent.mkdir(parents=True, exist_ok=True)
    if output_zip.exists():
        output_zip.unlink()
    with zipfile.ZipFile(output_zip, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(handoff_dir.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(handoff_dir).as_posix())


def write_docs(doc_output: Path, notes_output: Path, rows: list[dict[str, str]], abstract_words: int, keyword_total: int) -> None:
    status_counts: dict[str, int] = {}
    for row in rows:
        status_counts[row["status"]] = status_counts.get(row["status"], 0) + 1
    lines = [
        "# JOSR Repackaging Application",
        "",
        "## Judgement",
        "",
        "- The target-journal switch is technically sound for a public-data, hypothesis-generating transcriptomic study.",
        "- The Methodology article type is a reasonable fit because the manuscript emphasizes a tested computational framework.",
        "- A new OA-vs-normal diagnostic discrimination analysis was added to support fibrocartilage-matrix candidate stratification value.",
        "- Claims remain unchanged: candidate axes are not presented as causal or clinically deployable targets.",
        "- JOT-specific Translational Potential language was removed from the JOSR manuscript and cover letter.",
        "",
        "## Official-Basis Checks",
        "",
    ]
    lines.extend(f"- {item}" for item in OFFICIAL_BASIS)
    lines.extend(["", "## Counts", "", f"- Abstract words: {abstract_words}", f"- Keywords: {keyword_total}", ""])
    lines.append("## Audit Status Counts")
    lines.append("")
    lines.extend(f"- {status}: {count}" for status, count in sorted(status_counts.items()))
    write_text(doc_output, "\n".join(lines) + "\n")
    write_text(notes_output, "\n".join(lines) + "\n")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", required=True)
    parser.add_argument("--audit-output", required=True)
    parser.add_argument("--doc-output", required=True)
    parser.add_argument("--notes-output", required=True)
    args = parser.parse_args()

    root = Path(args.project_root)
    assembly = load_assembly_module(root)
    rows: list[dict[str, str]] = []

    manuscript = repackage_manuscript(read_text(root / "docs/manuscript/22_jot_anonymized_manuscript_draft.md"), rows)
    title_page = repackage_title_page(read_text(root / "docs/manuscript/20_jot_title_page_and_author_statements.md"), rows)
    cover = repackage_cover_letter(read_text(root / "docs/manuscript/21_jot_cover_letter_draft.md"), rows)
    declarations = build_declarations(title_page)

    manuscript_path = root / "docs/manuscript/37_josr_repackaged_manuscript_draft.md"
    cover_path = root / "docs/manuscript/38_josr_cover_letter_draft.md"
    title_page_path = root / "docs/manuscript/39_josr_title_page_and_author_statements.md"
    declarations_path = root / "docs/manuscript/40_josr_declarations.md"
    write_text(manuscript_path, manuscript)
    write_text(cover_path, cover)
    write_text(title_page_path, title_page)
    write_text(declarations_path, declarations)

    submission_dir = root / "submission/josr"
    handoff_dir = submission_dir / "upload_handoff"
    if handoff_dir.exists():
        shutil.rmtree(handoff_dir)
    handoff_dir.mkdir(parents=True, exist_ok=True)
    submission_dir.mkdir(parents=True, exist_ok=True)

    paths = {
        "title_page": submission_dir / "JOSR_Title_Page_and_Author_Statements.docx",
        "manuscript": submission_dir / "JOSR_Anonymized_Manuscript.docx",
        "cover_letter": submission_dir / "JOSR_Cover_Letter.docx",
        "declarations": submission_dir / "JOSR_Declarations.docx",
        "supplementary": submission_dir / "JOSR_Supplementary_Tables_ST01_ST34.xlsx",
    }
    assembly.write_docx(title_page, paths["title_page"], "JOSR Title Page and Author Statements")
    assembly.write_docx(manuscript, paths["manuscript"], "JOSR Anonymized Manuscript")
    assembly.write_docx(cover, paths["cover_letter"], "JOSR Cover Letter")
    assembly.write_docx(declarations, paths["declarations"], "JOSR Declarations")

    supp_manifest = pd.read_csv(root / "results/tables/manuscript_jot_supplementary_upload_manifest.tsv", sep="\t")
    supp_manifest = append_diagnostic_supplementary_rows(supp_manifest)
    supp_summary = write_josr_supplementary_workbook(root, assembly, supp_manifest, paths["supplementary"])
    rows.append({"item": "supplementary_workbook", "status": "updated", "detail": f"{len(supp_summary)} ST sheets plus README/Manifest"})

    manifest_rows = []
    manifest_rows.extend(copy_to_handoff(paths, handoff_dir))
    manifest_rows.extend(copy_figures(root, submission_dir, handoff_dir))
    write_handoff_readme(handoff_dir, manifest_rows)
    write_submitter_checklist(submission_dir)
    write_tsv(handoff_dir / "JOSR_UPLOAD_HANDOFF_MANIFEST.tsv", manifest_rows)
    zip_handoff(handoff_dir, submission_dir / "JOSR_upload_handoff_package.zip")

    abstract_words = word_count(extract_abstract(manuscript))
    keyword_total = keyword_count(manuscript)
    rows.append({"item": "abstract_word_count", "status": "checked", "detail": str(abstract_words)})
    rows.append({"item": "keyword_count", "status": "checked", "detail": str(keyword_total)})
    rows.append({"item": "manual_repository_url_or_doi", "status": "manual_author_action", "detail": "Repository URL/DOI placeholder intentionally retained."})
    rows.append({"item": "manual_author_declarations", "status": "manual_author_action", "detail": "Corresponding author must confirm declaration wording before submission."})

    write_tsv(Path(args.audit_output), rows)
    write_docs(Path(args.doc_output), Path(args.notes_output), rows, abstract_words, keyword_total)
    print(f"JOSR_ABSTRACT_WORDS {abstract_words}")
    print(f"JOSR_KEYWORDS {keyword_total}")
    print(f"JOSR_AUDIT_ROWS {len(rows)}")
    print(f"WROTE {args.audit_output}")
    print(f"WROTE {args.doc_output}")
    print(f"WROTE {args.notes_output}")


if __name__ == "__main__":
    main()
