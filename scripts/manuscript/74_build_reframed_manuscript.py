#!/usr/bin/env python
"""Create the reviewer-driven, structurally reframed manuscript."""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TITLE = (
    "Reproducible meniscal structural remodeling programs and context-dependent "
    "senescence-associated states in knee osteoarthritis: an integrative transcriptomic study"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, required=True)
    return parser.parse_args()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def set_margins(section) -> None:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)

    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(3)


def add_centered(document: Document, text: str, size: float = 10, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first, remainder = text.split(" ", 1)
        run = paragraph.add_run(first + " ")
        run.bold = True
        paragraph.add_run(remainder)
    else:
        paragraph.add_run(text)


def add_abstract_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(label + ": ")
    run.bold = True
    paragraph.add_run(text)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5)
        if widths:
            for index, width in enumerate(widths):
                cells[index].width = Cm(width)
    document.add_paragraph()


def add_caption(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(label + ". ")
    run.bold = True
    paragraph.add_run(text)


def add_reference(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.75)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(f"{number}. {text}")


def build_document(output: Path) -> None:
    document = Document()
    for section in document.sections:
        set_margins(section)
        add_page_number(section.footer.paragraphs[0])
    configure_styles(document)
    document.core_properties.title = TITLE
    document.core_properties.subject = "Knee osteoarthritis; meniscus; transcriptomic program discovery"
    document.core_properties.author = "Shiyou Ren; Dan Li; Ya Ding; Xilong Cui; Haiyang Yu"

    document.add_paragraph(TITLE, style="Title")
    add_centered(document, "Shiyou Ren1,3, Dan Li1,2, Ya Ding1,2, Xilong Cui1,2, Haiyang Yu1,2*", 11, True)
    add_centered(document, "1 Department of Orthopedics, Affiliated Fuyang People's Hospital of Anhui Medical University, Fuyang, Anhui Province, China", 9.5)
    add_centered(document, "2 National Key Clinical Specialty, Clinical Research Center for Spinal Deformity of Anhui Province, Fuyang, Anhui Province, China", 9.5)
    add_centered(document, "3 Department of Sports Medicine, The Eighth Affiliated Hospital of Sun Yat-sen University, Shenzhen, China", 9.5)
    add_centered(document, "* Corresponding author: Haiyang Yu, MD; Email: fy.yhy@163.com", 9.5)
    add_centered(document, "Short title: Meniscal remodeling programs in knee osteoarthritis", 9.5)
    add_centered(document, "Article type: Research Article", 9.5)

    document.add_heading("Abstract", level=1)
    add_abstract_paragraph(
        document,
        "Background",
        "Senescence-related interpretations of degenerative meniscus transcriptomes can be sensitive to sparse markers and investigator-defined scores. We asked which fibrochondrocyte programs emerge without forced signature genes, reproduce after independent decomposition of a second single-cell dataset, and show support in tissue-specific bulk cohorts.",
    )
    add_abstract_paragraph(
        document,
        "Methods",
        "We performed sample-balanced non-negative matrix factorization in 54,882 GSE220243 fibrochondrocyte/progenitor-like cells using 3,000 highly variable genes, four ranks, and five random starts. Rank-14 modules were tested across 6 osteoarthritis (OA) and 8 normal samples. HRA001986 was decomposed independently from 2,000 author-defined variable genes (28,686 cells; 6 abnormal and 6 normal samples) before post hoc module matching. Locked top-30 gene sets were then evaluated in six OA-vs-normal bulk cohorts using Hedges g and tissue-stratified random-effects models.",
    )
    add_abstract_paragraph(
        document,
        "Results",
        "GSE220243 contained an OA-enriched collagen/fibrotic remodeling module led by COL1A1, COL3A1, COL1A2, SPARC, TNC, and ASPN (OA-normal difference 0.601; permutation false discovery rate [FDR] 0.014). HRA001986 independently recovered a closely matched module (top-50 Jaccard 0.431) that was higher in abnormal samples (difference 0.730; FDR 0.022). A fibrocartilage-identity module also reproduced and was higher in abnormal HRA001986 samples, although its GSE220243 association did not pass FDR correction. Conversely, the senescence/stress module was lower in OA in GSE220243 (difference -0.227; FDR 0.028) and was not increased in HRA001986 (FDR 0.250). Both structural modules were positive in all three cartilage cohorts, but meniscus and synovium effects were inconsistent and tissue-specific meta-analytic intervals crossed zero.",
    )
    add_abstract_paragraph(
        document,
        "Conclusions",
        "Independent single-cell analyses support reproducible structural remodeling programs in degenerative meniscus, whereas the available data do not support a universal OA-upregulated meniscal senescence program. Bulk effects are tissue dependent. Both structural remodeling and senescence require protein-level, spatial, and functional validation.",
    )
    add_body(document, "Keywords: knee osteoarthritis; meniscus; fibrochondrocyte; single-cell RNA sequencing; non-negative matrix factorization; extracellular matrix; cellular senescence")

    document.add_heading("Introduction", level=1)
    add_body(document, "Knee osteoarthritis (OA) is a whole-joint disorder in which cartilage loss, subchondral bone change, synovitis, and periarticular tissue dysfunction evolve together [1]. The meniscus is central to load transmission and joint stability, and meniscal degeneration is common in established OA [2,3]. Meniscal injury and meniscectomy also increase the subsequent risk of knee OA [4]. These observations make the degenerative meniscus more than an incidental imaging or arthroscopic finding, but the cellular processes that distinguish adaptive fibrocartilage remodeling from pathologic change remain incompletely resolved.")
    add_body(document, "Human single-cell studies have described meniscal progenitor populations, inner- and outer-zone phenotypes, and degeneration-associated fibrochondrocyte states [5,6]. A separate study reported a senescent cell population in OA cartilage and meniscus and implicated ZEB1 as a regulator [8]. Cellular senescence is therefore biologically plausible in OA [7], but its transcriptomic identification is difficult. CDKN1A, CDKN2A, LMNB1, and secretory genes are sparse, state dependent, and not individually specific. Published gene sets such as Fridman, CellAge, the SASP atlas, and SenMayo capture overlapping but non-identical aspects of senescence [9-12]. A module enriched for these genes may represent durable senescence, transient stress, dissociation response, or a mixture of states.")
    add_body(document, "The original analysis underlying this work used curated signatures and a weighted composite score to prioritize a putative meniscus senescence program. That design could not establish whether the central signal was an independently identifiable biological program or a consequence of investigator-defined weights and forced gene inclusion. We therefore rebuilt the analysis around a stricter question: which transcriptional programs are discovered from highly variable genes alone, reproduce after independent decomposition of a second meniscus dataset, and show consistent sample-level disease associations?")
    add_body(document, "We hypothesized that structural extracellular-matrix programs would be more reproducible across datasets than senescence-associated states. To test this, we used sample-balanced non-negative matrix factorization (NMF) in GSE220243, independent NMF in HRA001986, and tissue-stratified evaluation in six OA-vs-normal bulk cohorts. Curated signatures were used only after discovery to annotate modules. The sample, rather than the cell, was the unit of disease-association testing, and cross-tissue pooling was treated as exploratory.")

    document.add_heading("Materials and methods", level=1)
    document.add_heading("Study design and data roles", level=2)
    add_body(document, "The analysis followed a discovery-reproduction-evaluation design (Table 1). GSE220243 was used for program discovery [8]. The author-processed HRA001986 meniscus h5ad object was used for independent program discovery and post hoc matching [6]. Six bulk OA-vs-normal cohorts were designated for primary evaluation: GSE114007, GSE143514, and GSE169077 (cartilage); GSE185064 (meniscus); and GSE55235 and GSE55457 (synovium). GSE89408 was retained as an RA-rich sensitivity cohort, while GSE98918 (OA versus arthroscopic partial meniscectomy) and GSE191157 (aged versus young) were contextual contrasts rather than primary OA-vs-normal evidence. No bulk cohort contributed to single-cell feature selection.")

    add_caption(document, "Table 1", "Data resources and designated analytical roles.")
    add_table(
        document,
        ["Resource", "Modality/tissue", "Samples", "Cells", "Analytical role"],
        [
            ["GSE220243", "Single-cell RNA-seq; meniscus", "14 (6 OA, 8 normal)", "54,882 selected fibrochondrocyte/progenitor-like cells", "HVG-only program discovery and sample-level testing"],
            ["HRA001986", "Single-cell RNA-seq; meniscus", "12 (6 abnormal, 6 normal)", "28,686 author-annotated chondrocyte-compartment cells", "Independent NMF; post hoc matching and sample-level reproduction"],
            ["GSE114007", "Bulk; cartilage", "38 (20 OA, 18 normal)", "NA", "Primary bulk evaluation"],
            ["GSE143514", "Bulk; cartilage", "8 (5 OA, 3 normal)", "NA", "Primary bulk evaluation; small-sample flag"],
            ["GSE169077", "Bulk; cartilage", "11 (6 OA, 5 normal)", "NA", "Primary bulk evaluation; small-sample flag"],
            ["GSE185064", "Bulk; meniscus", "8 (4 OA, 4 normal)", "NA", "Primary bulk evaluation; small-sample flag"],
            ["GSE55235", "Bulk; synovium", "20 (10 OA, 10 normal)", "NA", "Primary bulk evaluation"],
            ["GSE55457", "Bulk; synovium", "20 (10 OA, 10 normal)", "NA", "Primary bulk evaluation"],
            ["GSE89408; GSE98918; GSE191157", "Bulk; synovium or meniscus", "See S1 File", "NA", "Sensitivity/context only"],
        ],
        [2.4, 3.3, 2.8, 4.6, 5.0],
    )

    document.add_heading("Ethics statement", level=2)
    add_body(document, "This study was a secondary analysis of public and author-provided de-identified processed transcriptomic datasets. The authors did not recruit participants, contact participants, collect new human tissue, or receive identifiable information. The present secondary analysis was not submitted to the Ethics Committee of the Affiliated Fuyang People's Hospital of Anhui Medical University for separate approval because Article 32 of the Measures for Ethical Review of Life Science and Medical Research Involving Humans permits exemption for research using legally obtained public data or anonymized information when participant welfare, sensitive personal information, and commercial interests are not adversely affected [14]. No approval reference number was assigned. Ethics approval and consent for the original studies were reported by the source investigators.")

    document.add_heading("GSE220243 preprocessing and HVG-only program discovery", level=2)
    add_body(document, "GSE220243 10x matrices were processed with sample-aware quality control, broad-compartment annotation, and exclusion of immune, endothelial, mural/smooth-muscle, and other non-fibrochondrocyte compartments. The resulting discovery object contained 54,882 fibrochondrocyte and progenitor-like cells from 14 samples. The reanalysis used only the 3,000 genes marked as batch-aware highly variable genes in the existing object. No senescence, extracellular-matrix, communication, or other signature gene was added to the discovery universe.")
    add_body(document, "Library-size-normalized log expression was used for NMF. To limit domination by cell-rich samples, discovery was repeated in a strict balanced subsample of 300 cells per sample (4,200 cells) and a sensitivity subsample capped at 1,000 cells per sample (14,000 cells). NMF was fitted at ranks 10, 12, 14, and 16 with five random starts (seeds 11, 23, 37, 53, and 71). Components from the five starts at each rank were grouped by agglomerative clustering of cosine distance, and the median loading profile defined each consensus module. Rank 14 was retained for the main sample-level comparison to maintain continuity with the original analysis; matching to ranks 10, 12, and 16 served as cross-rank sensitivity analysis. NMF was used as an unsupervised program-discovery method [13].")
    add_body(document, "For each rank-14 consensus module, the 30 highest-loading genes were scored in all selected GSE220243 fibrochondrocytes after gene-wise standardization. Cell scores were averaged within each sample. OA-versus-normal differences were assessed across 6 OA and 8 normal samples using an exact label-permutation test; Benjamini-Hochberg FDR correction was applied across the 14 rank-14 modules [15]. Cliff's delta and exact Mann-Whitney tests were retained as descriptive effect summaries. Curated fibrocartilage, fibrotic-ECM, senescence, SASP, angiogenesis/hypoxia, paracrine, and generic-stress panels were tested against the top 50 genes of each module by hypergeometric enrichment only after NMF; these annotations did not influence program construction or sample scores.")

    document.add_heading("Independent HRA001986 program discovery and matching", level=2)
    add_body(document, "The HRA001986 analysis used the author-processed meniscal chondrocyte h5ad object containing 28,686 cells from 12 samples (6 abnormal and 6 normal, with inner- and outer-zone samples represented) [6]. Because the supplied object contained normalized non-negative expression but no raw-count layer, NMF was performed on the 2,000 genes identified as variable by the source authors. The same ranks, random starts, and sample-balanced subsampling schemes used for GSE220243 were applied independently, without adding GSE-derived genes to the HRA input.")
    add_body(document, "Only after HRA modules had been discovered were their top 50 genes compared with four GSE220243 reference modules: collagen/fibrotic remodeling, fibrocartilage identity, senescence/stress, and angiogenic/hypoxia. Similarity was measured by Jaccard overlap and a hypergeometric test across the shared variable-gene universe. Candidate matches were required to be represented in at least four of the five random starts. The top 30 genes of the best stable HRA match were scored across all HRA cells and averaged by sample. Abnormal-versus-normal differences were assessed by exact label permutation, with FDR correction across the four tested programs.")

    document.add_heading("Bulk cohort scoring and tissue-stratified evaluation", level=2)
    add_body(document, "For each of the four rank-14 GSE220243 programs, the 30 highest-loading genes were locked before bulk evaluation. Microarray probes were mapped to gene symbols and collapsed by mean expression; count or FPKM matrices were converted to gene-level matrices and log transformed when required. Within each cohort, each gene was standardized across samples, and the program score was the mean standardized expression of available genes. This cohort-internal procedure limits direct scale differences but cannot eliminate platform or tissue-composition effects.")
    add_body(document, "Within each designated OA-vs-normal cohort, standardized mean differences were expressed as Hedges g, with positive values indicating higher scores in OA. Random-effects meta-analysis was conducted within tissue strata using the Paule-Mandel estimator of between-study variance [16]. Because each tissue stratum contained few cohorts, uncertainty was calculated with a modified Hartung-Knapp procedure that used the larger of the conventional random-effects and Hartung-Knapp standard errors and a t distribution with k-1 degrees of freedom [17,18]. A single meniscus cohort was reported without pooling. An all-tissue model was retained only as an exploratory summary. FDR was controlled separately within each analytical scope. Directional area under the receiver operating characteristic curve and stratified bootstrap intervals [19,20] were generated as descriptive supplementary outputs only; no classifier was fitted, optimized, or claimed.")

    document.add_heading("Reporting hierarchy", level=2)
    add_body(document, "Evidence was interpreted in the following order: sample-level association in HVG-only GSE220243 discovery; gene-module reproduction and sample-level association after independent HRA001986 NMF; then tissue-specific bulk direction and uncertainty. A module was not considered an OA-associated senescence program merely because it contained CDKN1A, heat-shock, inflammatory, or SASP genes. No weighted senescence composite, mechanism evidence score, subtype classifier, pooled cross-tissue diagnostic AUC, or ligand-receptor mechanism was used for the central inference.")

    document.add_heading("Results", level=1)
    document.add_heading("HVG-only discovery prioritizes structural remodeling rather than an OA-upregulated senescence program", level=2)
    add_body(document, "HVG-only NMF separated structural, stress-related, and hypoxia-related transcriptional states without forcing signature genes into the input (Fig 1A-B; S1 File, Tables S2-S5). Cross-rank top-50 Jaccard similarity was 0.667 for the collagen/fibrotic module, 0.515 for fibrocartilage identity, 0.695 for senescence/stress, and 0.493 for angiogenic/hypoxia, indicating that all four patterns were recoverable at neighbouring ranks even though their disease directions differed.")
    add_body(document, "The strongest OA-positive rank-14 module was C12, led by COL1A1, COL3A1, COL1A2, SPARC, TNC, ASPN, DPT, COL5A2, AEBP1, and COL6A3. Its sample-mean score was higher in OA than normal meniscus (difference 0.601; Cliff's delta 1.00; exact permutation FDR 0.014). C6 contained canonical fibrocartilage genes including COMP, SCRG1, CILP, PRELP, CILP2, OGN, and MGP and was also higher in OA, but the association did not pass correction across all 14 modules (difference 0.313; FDR 0.078).")
    add_body(document, "The direction of senescence-related states was different. C5 contained CDKN1A, HMOX1, GADD45B, NFKBIA, ATF3, IER3, heat-shock genes, and immediate-early response genes. Although this module showed strong post hoc enrichment for senescence and generic stress panels, its score was lower in OA samples (difference -0.227; FDR 0.028). The angiogenic/hypoxia module C7, marked by ANGPTL4, VEGFA, ADM, BNIP3, and SLC2A1, was also lower in OA (difference -0.410; FDR 0.009). Thus, program reproducibility and OA direction were distinct properties (Fig 1C; S1 File, Tables S6-S8).")

    document.add_heading("Independent HRA001986 decomposition reproduces the structural programs", level=2)
    add_body(document, "Independent HRA001986 NMF recovered close counterparts of both structural programs under both cell-sampling schemes (Fig 2A; S1 File, Tables S9-S10). The best stable collagen/fibrotic match overlapped 28 of the top genes (Jaccard 0.431; overlap FDR 1.6 x 10^-36), while the fibrocartilage-identity match overlapped 27 genes (Jaccard 0.422; overlap FDR 2.2 x 10^-35). These matches were not produced by projecting GSE scores into HRA; the HRA modules were generated before comparison.")
    add_body(document, "At the sample level, the independent HRA collagen/fibrotic module was higher in abnormal than normal samples (difference 0.730; permutation FDR 0.022), as was the fibrocartilage-identity module (difference 0.594; FDR 0.024). A senescence/stress-like HRA module was identifiable by gene overlap (22 top genes; Jaccard 0.324; overlap FDR 5.5 x 10^-26), but its abnormal-normal difference was not significant (difference 0.235; FDR 0.250). The angiogenic/hypoxia match was weaker (Jaccard 0.164) and was also not disease associated (FDR 0.250) (Fig 2B; S1 File, Tables S11-S12).")

    document.add_heading("Bulk effects are strongest in cartilage and do not support a universal cross-tissue signature", level=2)
    add_body(document, "The locked collagen/fibrotic program was higher in OA in all three cartilage cohorts (Hedges g 2.29, 0.21, and 1.79), but the tissue-stratified random-effects estimate remained imprecise (g 1.48; modified Hartung-Knapp 95% CI -1.22 to 4.19; I2 73%). The two synovium cohorts were discordant (g 2.00 and -0.13), and the single meniscus cohort was negative but uncertain (g -0.33; 95% CI -1.55 to 0.89) (Fig 3A).")
    add_body(document, "The fibrocartilage-identity program was likewise positive in all cartilage cohorts (tissue-stratified g 0.95; 95% CI -0.23 to 2.12; I2 0%) but negative in the meniscus cohort and one synovium cohort. Across the six primary cohorts, four effects were positive and two negative for each structural program. Exploratory all-tissue estimates crossed zero for both collagen/fibrotic remodeling (g 0.98; 95% CI -0.27 to 2.23; FDR 0.115) and fibrocartilage identity (g 0.55; 95% CI -0.19 to 1.29; FDR 0.115). These results support tissue-dependent structural remodeling, not a universal cross-tissue biomarker.")
    add_body(document, "The senescence/stress score was lower in OA in five of six primary cohorts, and the angiogenic/hypoxia score was lower in all six (Fig 3B). Their exploratory all-tissue estimates were negative, but neither passed FDR correction within the all-tissue scope (senescence/stress g -1.67, 95% CI -2.89 to -0.45, FDR 0.067; angiogenic/hypoxia g -1.77, 95% CI -3.44 to -0.09, FDR 0.085). Tissue-specific intervals were wide. These findings reinforce the absence of a consistently OA-upregulated senescence or angiogenic program in the available transcriptomic cohorts (Table 2; S1 File, Tables S13-S17).")

    document.add_page_break()
    add_caption(document, "Table 2", "Evidence hierarchy for the four HVG-only reference programs.")
    add_table(
        document,
        ["Program", "GSE220243", "Independent HRA001986", "Primary bulk cohorts", "Interpretation"],
        [
            ["Collagen/fibrotic remodeling", "OA-normal delta 0.601; FDR 0.014", "Jaccard 0.431; abnormal-normal delta 0.730; FDR 0.022", "4 positive, 2 negative; cartilage 3/3 positive", "Primary reproducible structural program; bulk association is tissue dependent"],
            ["Fibrocartilage identity", "Delta 0.313; FDR 0.078", "Jaccard 0.422; delta 0.594; FDR 0.024", "4 positive, 2 negative; cartilage 3/3 positive", "Supportive structural program"],
            ["Senescence/stress", "Delta -0.227; FDR 0.028", "Jaccard 0.324; delta 0.235; FDR 0.250", "1 positive, 5 negative", "Reproducible state, but not a universal OA-upregulated program"],
            ["Angiogenic/hypoxia", "Delta -0.410; FDR 0.009", "Jaccard 0.164; delta 0.135; FDR 0.250", "0 positive, 6 negative", "Exploratory context state; no OA-up evidence"],
        ],
        [3.1, 3.2, 4.4, 3.2, 5.0],
    )

    document.add_heading("Discussion", level=1)
    document.add_heading("Principal findings", level=2)
    add_body(document, "This reanalysis answers a narrower and more defensible question than the original senescence-centered framework. When programs were discovered from highly variable genes alone and disease testing was performed at the sample level, the clearest reproducible signal was structural remodeling. A collagen-rich fibrotic program was increased in OA meniscus in GSE220243, independently recovered in HRA001986, and higher in abnormal HRA samples. A second fibrocartilage-identity program also reproduced across the two single-cell datasets. These are the principal positive findings.")
    add_body(document, "The senescence result was more qualified. A CDKN1A- and stress-enriched module was reproducible as a transcriptional state, but it was lower in OA in GSE220243, not significantly increased in HRA001986, and lower in five of six primary bulk cohorts. The data therefore do not support a single, universally OA-upregulated meniscal senescence program. This distinction matters because reproducibility of a gene module does not establish its disease direction, biological permanence, or causal role.")

    document.add_heading("Structural programs are robust at the single-cell level but tissue dependent in bulk", level=2)
    add_body(document, "The independent HRA analysis addresses an important limitation of score projection. The second dataset was decomposed without GSE-derived genes, and reference modules were introduced only after discovery. The strong overlaps for collagen/fibrotic remodeling and fibrocartilage identity therefore indicate genuine recurrence of gene co-expression structure rather than forced score transfer. The abnormal-normal differences in HRA further support disease relevance, although each group contained only six samples.")
    add_body(document, "Bulk evaluation was less uniform. Both structural programs were positive in all cartilage cohorts, but not in the meniscus or synovium cohorts. Several explanations are plausible: bulk tissues contain different cell mixtures; the meniscus cohort was small; disease stage and comparator definitions varied; and a meniscus-derived cellular program may not preserve the same direction in cartilage or synovium. The wide modified Hartung-Knapp intervals appropriately reflect the small number of cohorts per tissue. We therefore regard the bulk analysis as evidence of context dependence rather than external confirmation of a common cross-tissue effect.")

    document.add_heading("Why a senescence/stress module should not be equated with cellular senescence", level=2)
    add_body(document, "The senescence/stress module included CDKN1A and HMOX1, but also heat-shock proteins, immediate-early genes, and stress-response regulators. Such a profile can arise from durable senescence, transient stress, tissue handling, or dissociation. The lack of a consistent disease-positive direction argues against naming this module a disease-driving senescence program on transcriptomic evidence alone. Multi-marker protein assessment, loss of proliferative capacity, senescence-associated beta-galactosidase, chromatin changes, and a context-appropriate secretory phenotype are needed to separate these possibilities [7,9-12].")
    add_body(document, "This interpretation does not exclude a role for senescent cells in OA or contradict experimental senescence studies [7,8]. It shows that the public datasets analysed here do not justify a stronger meniscus-specific claim. The negative or null disease-direction results are informative because they identify where transcriptomic annotation outruns the evidence.")

    document.add_heading("Clinical and translational implications", level=2)
    add_body(document, "For orthopaedic and sports-medicine research, the revised findings favour direct validation of structural remodeling at the meniscus surface, inner and outer zones, and meniscus-synovium interface. COL1A1, COL3A1, SPARC, ASPN, TNC, COMP, CILP, SCRG1, FMOD, and related proteins provide a focused panel for histology, immunostaining, or spatial assays. Clinical correlations should examine tear pattern, meniscal extrusion, cartilage status, synovitis, and OA stage rather than assuming that one molecular score applies across all joint tissues.")
    add_body(document, "Senescence remains a reasonable secondary hypothesis, but future studies should define it phenotypically rather than by a single transcript or a weighted bioinformatic score. Prospective tissue collection with age, sex, body mass index, mechanical alignment, and imaging metadata would also permit separation of normal aging, traumatic injury, and OA-associated remodeling.")

    document.add_heading("Strengths and limitations", level=2)
    add_body(document, "The main strengths are removal of forced signature genes from discovery, repeated NMF across ranks and random starts, sample-balanced cell selection, sample-level rather than cell-level disease testing, independent decomposition of HRA001986, explicit separation of data roles, and conservative tissue-stratified meta-analysis. The full direction of non-supportive results is reported.")
    add_body(document, "Several limitations remain. First, both single-cell datasets are observational and small at the participant/sample level. Second, the HRA object was author processed and lacked a raw-count layer, so its NMF input differed from GSE220243. Third, module matching and top-30 scoring require analytical choices, even though these choices no longer determine the discovery gene universe. Fourth, the primary bulk cohorts differed in platform, tissue, disease stage, and cell composition; three cohorts had fewer than six samples in at least one group. Fifth, age, sex, body mass index, alignment, and OA grade were not consistently available for covariate adjustment. Sixth, no spatial, protein-level, proliferative, or perturbation experiment was performed. Finally, the study does not establish a diagnostic classifier, causal intercellular pathway, or treatment target.")

    document.add_heading("Conclusions", level=1)
    add_body(document, "HVG-only program discovery and independent HRA001986 decomposition identify reproducible collagen/fibrotic-remodeling and fibrocartilage-identity programs in human meniscal cells. The collagen/fibrotic program was increased in OA or abnormal meniscus in both single-cell datasets, while bulk effects varied by tissue and cohort. A senescence/stress state was reproducible at the gene-module level but was not consistently increased in OA. These data support structural remodeling as the primary transcriptomic finding and place senescence as a context-dependent hypothesis requiring direct phenotypic validation.")

    document.add_heading("Abbreviations", level=1)
    add_body(document, "AUC, area under the receiver operating characteristic curve; CI, confidence interval; ECM, extracellular matrix; FDR, false discovery rate; HRA, GSA-Human Run Archive; HVG, highly variable gene; NMF, non-negative matrix factorization; OA, osteoarthritis; RA, rheumatoid arthritis; SASP, senescence-associated secretory phenotype.")

    document.add_heading("Data availability", level=1)
    add_body(document, "Public source data are available from the Gene Expression Omnibus under GSE220243, GSE114007, GSE143514, GSE169077, GSE185064, GSE191157, GSE55235, GSE55457, GSE89408, and GSE98918, and from GSA-Human under HRA001986 (https://ngdc.cncb.ac.cn/gsa-human/browse/HRA001986). Analysis code and derived, non-identifiable outputs are available at https://github.com/shiyouren91/meniscus-senescence-program-koa. The author-provided processed HRA001986 h5ad files are not redistributed; the underlying sequencing data are available through HRA001986, and requests for the exact processed files may be directed to the source-study authors. Revised supplementary tables provide the data points underlying the reported module matches, sample-level effects, cohort effects, and meta-analytic summaries.")

    document.add_heading("Ethics approval and consent to participate", level=1)
    add_body(document, "Not applicable to this secondary analysis for the reasons stated in the Ethics statement. Ethics approval and consent for the original human studies were the responsibility of the source investigators.")

    document.add_heading("Consent for publication", level=1)
    add_body(document, "Not applicable.")

    document.add_heading("Competing interests", level=1)
    add_body(document, "The authors declare that they have no competing interests.")

    document.add_heading("Funding", level=1)
    add_body(document, "This research was funded by the Clinical Medicine Translational Research Special Program of Anhui Provincial Department of Science and Technology (grant number 202527c10020008).")

    document.add_heading("Author contributions", level=1)
    add_body(document, "H.Y. and S.R. conceived and designed the study. S.R. performed the data analysis and wrote the original draft. D.L. and Y.D. assisted with data curation and visualization. X.C. and D.L. contributed to methodology and result interpretation. H.Y. supervised the project, acquired funding, and revised the manuscript. All authors read and approved the final manuscript.")

    document.add_heading("Acknowledgments and use of artificial intelligence", level=1)
    add_body(document, "OpenAI Codex was used to assist with code drafting, quality-control checks, language editing, document organization, and formatting. All scripts, outputs, statistical decisions, and interpretations were reviewed by the authors, who take full responsibility for the work.")

    document.add_page_break()
    document.add_heading("References", level=1)
    references = [
        "Katz JN, Arant KR, Loeser RF. Diagnosis and treatment of hip and knee osteoarthritis: a review. JAMA. 2021;325(6):568-578. doi:10.1001/jama.2020.22171.",
        "Ozeki N, Koga H, Sekiya I. Degenerative meniscus in knee osteoarthritis: from pathology to treatment. Life (Basel). 2022;12(4):603. doi:10.3390/life12040603.",
        "Ma Z, Vyhlidal MJ, Li DX, Adesida AB. Mechano-bioengineering of the knee meniscus. Am J Physiol Cell Physiol. 2022;323(6):C1652-C1663. doi:10.1152/ajpcell.00336.2022.",
        "Whittaker JL, Losciale JM, Juhl CB, Thorlund JB, Lundberg M, Truong LK, et al. Risk factors for knee osteoarthritis after traumatic knee injury: a systematic review and meta-analysis of randomised controlled trials and cohort studies for the OPTIKNEE Consensus. Br J Sports Med. 2022;56(24):1406-1421. doi:10.1136/bjsports-2022-105496.",
        "Sun H, Wen X, Li H, Wu P, Gu M, Zhao X, et al. Single-cell RNA-seq analysis identifies meniscus progenitors and reveals the progression of meniscus degeneration. Ann Rheum Dis. 2020;79(3):408-417. doi:10.1136/annrheumdis-2019-215926.",
        "Fu W, Chen S, Yang R, Li C, Gao H, Li J, et al. Cellular features of localized microenvironments in human meniscal degeneration: a single-cell transcriptomic study. eLife. 2022;11:e79585. doi:10.7554/eLife.79585.",
        "Coryell PR, Diekman BO, Loeser RF. Mechanisms and therapeutic implications of cellular senescence in osteoarthritis. Nat Rev Rheumatol. 2021;17(1):47-57. doi:10.1038/s41584-020-00533-7.",
        "Swahn H, Li K, Duffy T, Olmer M, D'Lima DD, Mondala TS, et al. Senescent cell population with ZEB1 transcription factor as its main regulator promotes osteoarthritis in cartilage and meniscus. Ann Rheum Dis. 2023;82(3):403-415. doi:10.1136/ard-2022-223227.",
        "Fridman AL, Tainsky MA. Critical pathways in cellular senescence and immortalization revealed by gene expression profiling. Oncogene. 2008;27(46):5975-5987. doi:10.1038/onc.2008.213.",
        "Avelar RA, Ortega JG, Tacutu R, Tyler EJ, Bennett D, Binetti P, et al. A multidimensional systems biology analysis of cellular senescence in aging and disease. Genome Biol. 2020;21(1):91. doi:10.1186/s13059-020-01990-9.",
        "Basisty N, Kale A, Jeon OH, Kuehnemann C, Payne T, Rao C, et al. A proteomic atlas of senescence-associated secretomes for aging biomarker development. PLoS Biol. 2020;18(1):e3000599. doi:10.1371/journal.pbio.3000599.",
        "Saul D, Kosinsky RL, Atkinson EJ, Doolittle ML, Zhang X, LeBrasseur NK, et al. A new gene set identifies senescent cells and predicts senescence-associated pathways across tissues. Nat Commun. 2022;13(1):4827. doi:10.1038/s41467-022-32552-1.",
        "Kotliar D, Veres A, Nagy MA, Tabrizi S, Hodis E, Melton DA, et al. Identifying gene expression programs of cell-type identity and cellular activity with single-cell RNA-seq. eLife. 2019;8:e43803. doi:10.7554/eLife.43803.",
        "National Health Commission of the People's Republic of China, Ministry of Education of the People's Republic of China, Ministry of Science and Technology of the People's Republic of China, National Administration of Traditional Chinese Medicine. Measures for Ethical Review of Life Science and Medical Research Involving Humans [Internet]. Beijing: National Health Commission of the People's Republic of China; 2023 [cited 2026 Jun 29]. Available from: https://www.nhc.gov.cn/qjjys/c100016/202302/6b6e447b3edc4338856c9a652a85f44b.shtml.",
        "Benjamini Y, Hochberg Y. Controlling the false discovery rate: a practical and powerful approach to multiple testing. J R Stat Soc Series B Stat Methodol. 1995;57(1):289-300. doi:10.1111/j.2517-6161.1995.tb02031.x.",
        "Paule RC, Mandel J. Consensus values and weighting factors. J Res Natl Bur Stand. 1982;87(5):377-385. doi:10.6028/JRES.087.022.",
        "Hartung J, Knapp G. On tests of the overall treatment effect in meta-analysis with normally distributed responses. Stat Med. 2001;20(12):1771-1782. doi:10.1002/sim.791.",
        "Röver C, Knapp G, Friede T. Hartung-Knapp-Sidik-Jonkman approach and its modification for random-effects meta-analysis with few studies. BMC Med Res Methodol. 2015;15:99. doi:10.1186/s12874-015-0091-1.",
        "Hanley JA, McNeil BJ. The meaning and use of the area under a receiver operating characteristic (ROC) curve. Radiology. 1982;143(1):29-36. doi:10.1148/radiology.143.1.7063747.",
        "Efron B. Bootstrap methods: another look at the jackknife. Ann Stat. 1979;7(1):1-26. doi:10.1214/aos/1176344552.",
    ]
    for number, reference in enumerate(references, start=1):
        add_reference(document, number, reference)

    document.add_heading("Figure legends", level=1)
    add_caption(document, "Fig 1", "HVG-only, sample-aware program discovery in GSE220243. (A) Programs were discovered from 3,000 batch-aware highly variable genes without forced inclusion of senescence, matrix, or communication genes. NMF was repeated across four ranks and five random starts, and disease associations were tested across 14 independent samples. (B) Curated signatures were used only after discovery to annotate rank-14 consensus modules; colour represents -log10(FDR) for top-50 gene overlap. (C) OA-minus-normal differences in sample-mean top-30 module scores. Blue bars indicate exact label-permutation FDR < 0.05 across the 14 modules.")
    add_caption(document, "Fig 2", "Independent program reproduction in HRA001986. (A) Median best top-50 Jaccard overlap across 20 independent HRA NMF runs for each of two sample-balanced cell selections. Error bars show the interquartile range. GSE220243 reference genes were introduced only after HRA program discovery. (B) Sample-mean scores for the best stable HRA match in six normal and six abnormal samples. Collagen/fibrotic remodeling and fibrocartilage identity were higher in abnormal samples after FDR correction; senescence/stress and angiogenic/hypoxia were not.")
    add_caption(document, "Fig 3", "Tissue-stratified bulk evaluation of locked GSE220243 programs. (A) Cohort-specific Hedges g and 95% confidence intervals for the collagen/fibrotic remodeling program. Positive values indicate higher scores in OA. (B) Hedges g for all four programs across the six designated OA-vs-normal cohorts. The structural programs were positive in all cartilage cohorts but not consistently positive in meniscus or synovium. Senescence/stress and angiogenic/hypoxia were generally lower in OA. Cross-tissue pooling was exploratory.")

    document.add_heading("Supporting information caption", level=1)
    add_caption(document, "S1 File", "Revised supplementary tables. The workbook contains the dataset-role audit, HVG-only GSE220243 NMF diagnostics and consensus programs, signature annotations, sample scores and exact permutation tests, cross-rank stability, independent HRA001986 matching and sample-level tests, bulk dataset and gene-coverage audits, per-cohort effects, tissue-stratified meta-analysis, and the final evidence hierarchy.")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def copy_figures(project_root: Path, output_dir: Path) -> None:
    source = project_root / "results" / "figures" / "reviewer_reanalysis" / "revised_main"
    destination = output_dir / "Figures"
    destination.mkdir(parents=True, exist_ok=True)
    for index, stem in enumerate(
        ["Fig1_HVG_only_program_discovery", "Fig2_HRA_independent_reproduction", "Fig3_tissue_stratified_bulk_validation"],
        start=1,
    ):
        shutil.copy2(source / f"{stem}.tif", destination / f"Fig{index}.tif")
        shutil.copy2(source / f"{stem}.png", destination / f"Fig{index}_preview.png")


def main() -> None:
    args = parse_args()
    root = args.project_root.resolve()
    output_dir = root / "submission" / "reviewer_reanalysis_20260925"
    manuscript = output_dir / "Revised_Manuscript_Structural_Remodeling_Primary.docx"
    build_document(manuscript)
    copy_figures(root, output_dir)
    print(manuscript)


if __name__ == "__main__":
    main()
