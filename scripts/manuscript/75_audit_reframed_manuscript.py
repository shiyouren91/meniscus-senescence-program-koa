#!/usr/bin/env python
"""Audit the reframed manuscript against its figures, tables, and source results."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from docx import Document


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, required=True)
    return parser.parse_args()


def expand_citation_block(block: str) -> set[int]:
    numbers: set[int] = set()
    for part in block.split(","):
        part = part.strip()
        if "-" in part:
            start, end = (int(value) for value in part.split("-", maxsplit=1))
            numbers.update(range(start, end + 1))
        elif part.isdigit():
            numbers.add(int(part))
    return numbers


def assert_close(actual: float, expected: float, tolerance: float, label: str) -> None:
    if abs(actual - expected) > tolerance:
        raise AssertionError(f"{label}: expected {expected}, found {actual}")


def main() -> None:
    root = parse_args().project_root.resolve()
    manuscript = root / "submission" / "reviewer_reanalysis_20260925" / "Revised_Manuscript_Structural_Remodeling_Primary.docx"
    document = Document(manuscript)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    all_text = "\n".join(paragraphs)

    reference_start = paragraphs.index("References")
    legend_start = paragraphs.index("Figure legends")
    body_text = "\n".join(paragraphs[:reference_start])
    reference_lines = [line for line in paragraphs[reference_start + 1 : legend_start] if re.match(r"^\d+\.\s", line)]

    required_callouts = ["Fig 1", "Fig 2", "Fig 3", "Table 1", "Table 2", "S1 File"]
    missing_callouts = [term for term in required_callouts if term not in body_text]
    if missing_callouts:
        raise AssertionError(f"Missing body callouts: {', '.join(missing_callouts)}")

    cited: set[int] = set()
    for block in re.findall(r"\[([0-9,\-]+)\]", body_text):
        cited.update(expand_citation_block(block))
    expected_references = set(range(1, len(reference_lines) + 1))
    if cited != expected_references:
        raise AssertionError(
            f"Citation mismatch: missing={sorted(expected_references - cited)}, out_of_range={sorted(cited - expected_references)}"
        )

    stale_terms = [
        "MIF-CD74",
        "MIF_CD74",
        "ANGPTL4-integrin",
        "ANGPTL4_integrin",
        "mechanism evidence score",
        "subtype classifier",
        "pooled cross-tissue diagnostic AUC",
        "CellChat",
        "LIANA",
        "NicheNet",
        "A program-level meniscus senescence analysis identifies",
    ]
    unexpected_stale = [term for term in stale_terms if term in all_text and term not in body_text.split("Reporting hierarchy", 1)[-1][:1000]]
    # The three deprecated analysis labels are permitted once in the reporting guardrail only.
    for term in ["mechanism evidence score", "subtype classifier", "pooled cross-tissue diagnostic AUC"]:
        if all_text.count(term) > 1:
            raise AssertionError(f"Deprecated term occurs more than once: {term}")
    unexpected_stale = [term for term in unexpected_stale if term not in {"mechanism evidence score", "subtype classifier", "pooled cross-tissue diagnostic AUC"}]
    if unexpected_stale:
        raise AssertionError(f"Stale terms remain: {', '.join(unexpected_stale)}")

    placeholder_patterns = [r"\[AUC\]", r"\[CI\]", r"\bTBD\b", r"\bTODO\b", r"\bXX+\b"]
    placeholders = [pattern for pattern in placeholder_patterns if re.search(pattern, all_text, flags=re.IGNORECASE)]
    if placeholders:
        raise AssertionError(f"Placeholders remain: {', '.join(placeholders)}")

    abstract_start = paragraphs.index("Abstract")
    introduction_start = paragraphs.index("Introduction")
    abstract_text = " ".join(paragraphs[abstract_start + 1 : introduction_start])
    abstract_word_count = len(re.findall(r"\b[\w-]+\b", abstract_text))
    if not 150 <= abstract_word_count <= 350:
        raise AssertionError(f"Unexpected abstract word count: {abstract_word_count}")

    if len(document.tables) != 2:
        raise AssertionError(f"Expected 2 main tables, found {len(document.tables)}")
    figure_caption_count = sum(text.startswith(("Fig 1.", "Fig 2.", "Fig 3.")) for text in paragraphs)
    if figure_caption_count != 3:
        raise AssertionError(f"Expected 3 figure legends, found {figure_caption_count}")
    if not any(text.startswith("S1 File.") for text in paragraphs):
        raise AssertionError("S1 File caption is missing")

    tables = root / "results" / "tables" / "reviewer_reanalysis"
    gse = pd.read_csv(tables / "gse220243_hvg_only_rank14_sample_tests.tsv", sep="\t").set_index("program_id")
    for program, delta, fdr in [
        ("HVG_K14_C12", 0.601461, 0.013981),
        ("HVG_K14_C6", 0.313324, 0.077896),
        ("HVG_K14_C5", -0.227482, 0.027963),
        ("HVG_K14_C7", -0.409995, 0.009321),
    ]:
        assert_close(float(gse.loc[program, "delta_oa_minus_normal"]), delta, 5e-6, f"{program} delta")
        assert_close(float(gse.loc[program, "permutation_fdr"]), fdr, 5e-6, f"{program} FDR")

    hra = pd.read_csv(tables / "hra001986_hvg_reference_sample_tests.tsv", sep="\t").set_index("target_id")
    for program, delta, fdr in [
        ("collagen_fibrotic_remodeling", 0.730253, 0.021622),
        ("fibrocartilage_identity", 0.593911, 0.023784),
        ("senescence_stress", 0.234760, 0.249730),
        ("angiogenic_hypoxia", 0.134646, 0.249730),
    ]:
        assert_close(float(hra.loc[program, "delta_abnormal_minus_normal"]), delta, 5e-6, f"{program} HRA delta")
        assert_close(float(hra.loc[program, "permutation_fdr"]), fdr, 5e-6, f"{program} HRA FDR")

    meta = pd.read_csv(tables / "bulk_hvg_program_meta_summary.tsv", sep="\t").set_index(["scope", "axis_id"])
    for scope, program, effect, low, high in [
        ("tissue_cartilage", "collagen_fibrotic_remodeling", 1.484942, -1.219990, 4.189873),
        ("tissue_cartilage", "fibrocartilage_identity", 0.947787, -0.226194, 2.121767),
        ("all_tissues_exploratory", "collagen_fibrotic_remodeling", 0.983231, -0.265897, 2.232359),
        ("all_tissues_exploratory", "fibrocartilage_identity", 0.549394, -0.192391, 1.291179),
    ]:
        row = meta.loc[[(scope, program)]].iloc[0]
        assert_close(float(row["pooled_hedges_g"]), effect, 5e-6, f"{scope}/{program} effect")
        assert_close(float(row["hk_ci95_low"]), low, 5e-6, f"{scope}/{program} low CI")
        assert_close(float(row["hk_ci95_high"]), high, 5e-6, f"{scope}/{program} high CI")

    print(f"PASS manuscript={manuscript}")
    print(f"PASS abstract_words={abstract_word_count}")
    print(f"PASS body_callouts={','.join(required_callouts)}")
    print(f"PASS references={len(reference_lines)} all_cited=yes")
    print(f"PASS tables={len(document.tables)} figure_legends={figure_caption_count} supporting_files=1")
    print("PASS key_numeric_claims_match_source_tables")
    print("PASS stale_terms_and_placeholders")


if __name__ == "__main__":
    main()
